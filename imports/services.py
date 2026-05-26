from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook

from catalog.models import Brand, Category, Product, ProductStock, Supplier, Warehouse


BASE_COLUMNS = [
    "codigo",
    "codigo_barra",
    "descripcion",
    "marca",
    "modelo",
    "categoria",
    "proveedor",
    "costo",
    "precio",
    "aplica_iva",
    "garantia",
    "observaciones",
]

REQUIRED_COLUMNS = {"codigo", "descripcion", "categoria", "costo", "precio"}


def normalize_header(value):
    return str(value or "").strip().lower().replace(" ", "_")


def normalize_text(value):
    return str(value or "").strip()


def parse_decimal(value):
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        return None


def parse_bool(value):
    text = normalize_text(value).lower()
    if text in {"si", "sí", "s", "1", "true", "verdadero", "x"}:
        return True
    if text in {"no", "n", "0", "false", "falso"}:
        return False
    return True


def build_template_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Productos"
    warehouse_columns = [
        f"stock_{warehouse.name.strip().lower().replace(' ', '_')}"
        for warehouse in Warehouse.objects.filter(is_active=True).order_by("name")
    ]
    headers = BASE_COLUMNS + warehouse_columns
    sheet.append(headers)
    sheet.append(
        [
            "PROD-001",
            "123456789",
            "Producto de ejemplo",
            "Marca",
            "Modelo",
            "Categoria",
            "Proveedor",
            100,
            150,
            "si",
            "30 dias",
            "Observacion opcional",
        ]
        + [0 for _ in warehouse_columns]
    )
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 20
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def preview_product_import(file):
    workbook = load_workbook(file, data_only=True)
    sheet = workbook.active
    headers = [normalize_header(cell.value) for cell in sheet[1]]
    missing = sorted(REQUIRED_COLUMNS - set(headers))
    rows = []
    duplicate_codes = set()
    seen_codes = set()

    warehouse_by_column = {}
    for header in headers:
        if header.startswith("stock_"):
            raw_name = header.replace("stock_", "").replace("_", " ").strip()
            if raw_name:
                warehouse_by_column[header] = raw_name

    if missing:
        return {
            "rows": [],
            "has_errors": True,
            "general_errors": ["Faltan columnas obligatorias: " + ", ".join(missing)],
        }

    for display_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=1):
        raw = dict(zip(headers, values))
        if not any(value not in (None, "") for value in raw.values()):
            continue

        code = normalize_text(raw.get("codigo"))
        errors = []
        warnings = []

        if not code:
            errors.append("Codigo obligatorio vacio.")
        elif code in seen_codes:
            errors.append("Codigo duplicado dentro del Excel.")
            duplicate_codes.add(code)
        else:
            seen_codes.add(code)

        description = normalize_text(raw.get("descripcion"))
        category = normalize_text(raw.get("categoria"))
        cost = parse_decimal(raw.get("costo"))
        price = parse_decimal(raw.get("precio"))

        if not description:
            errors.append("Descripcion obligatoria vacia.")
        if not category:
            errors.append("Categoria obligatoria vacia.")
        if cost is None or cost < 0:
            errors.append("Costo invalido.")
        if price is None or price < 0:
            errors.append("Precio invalido.")

        exists = Product.objects.filter(code=code).exists() if code else False
        if exists:
            warnings.append("El codigo ya existe en el sistema.")

        stock_values = {}
        for column, warehouse_name in warehouse_by_column.items():
            quantity = parse_decimal(raw.get(column))
            if quantity is None or quantity < 0:
                errors.append(f"Stock invalido en {column}.")
            else:
                stock_values[warehouse_name] = str(quantity)

        rows.append(
            {
                "display_row": display_row,
                "code": code,
                "description": description,
                "brand": normalize_text(raw.get("marca")),
                "model": normalize_text(raw.get("modelo")),
                "category": category,
                "supplier": normalize_text(raw.get("proveedor")),
                "barcode": normalize_text(raw.get("codigo_barra")),
                "cost": str(cost if cost is not None else ""),
                "price": str(price if price is not None else ""),
                "applies_iva": parse_bool(raw.get("aplica_iva")),
                "warranty": normalize_text(raw.get("garantia")),
                "notes": normalize_text(raw.get("observaciones")),
                "stock": stock_values,
                "exists": exists,
                "errors": errors,
                "warnings": warnings,
            }
        )

    return {
        "rows": rows,
        "has_errors": any(row["errors"] for row in rows),
        "general_errors": [],
        "duplicate_codes": sorted(duplicate_codes),
    }


def apply_product_import(rows, duplicate_action):
    created = 0
    updated = 0
    ignored = 0

    for row in rows:
        if row["errors"]:
            ignored += 1
            continue

        product = Product.objects.filter(code=row["code"]).first()
        if product and duplicate_action == "ignore":
            ignored += 1
            continue

        category, _ = Category.objects.get_or_create(
            name=row["category"],
            defaults={"is_active": True},
        )
        brand = None
        if row["brand"]:
            brand, _ = Brand.objects.get_or_create(name=row["brand"])
        supplier = None
        if row["supplier"]:
            supplier, _ = Supplier.objects.get_or_create(name=row["supplier"])

        data = {
            "barcode": row["barcode"],
            "description": row["description"],
            "brand": brand,
            "model": row["model"],
            "category": category,
            "supplier": supplier,
            "cost_price": Decimal(row["cost"]),
            "sale_price": Decimal(row["price"]),
            "applies_iva": row["applies_iva"],
            "warranty": row["warranty"],
            "notes": row["notes"],
            "is_active": True,
        }

        if product:
            for field, value in data.items():
                setattr(product, field, value)
            product.save()
            updated += 1
        else:
            product = Product.objects.create(code=row["code"], **data)
            created += 1

        for warehouse_name, quantity in row["stock"].items():
            warehouse, _ = Warehouse.objects.get_or_create(
                name=warehouse_name.title(),
                defaults={"is_active": True},
            )
            ProductStock.objects.update_or_create(
                product=product,
                warehouse=warehouse,
                defaults={"quantity": Decimal(quantity)},
            )

    return {"created": created, "updated": updated, "ignored": ignored}
